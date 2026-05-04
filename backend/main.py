"""VocabForge — FastAPI backend server."""

import os
import tempfile

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from processor import process_dialogues

app = FastAPI(title="VocabForge")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class ProcessRequest(BaseModel):
    dialogues: str
    vocab: list[str] = []
    title: str = "学习笔记"


@app.post("/api/process")
def process(req: ProcessRequest):
    return process_dialogues(req.dialogues, set(req.vocab), req.title)


def _extract_text(content: bytes, ext: str) -> str:
    """Write content to temp file, parse with appropriate library, return text."""
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
    try:
        tmp.write(content)
    finally:
        tmp.close()

    try:
        if ext == '.pdf':
            import pdfplumber
            with pdfplumber.open(tmp.name) as pdf:
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)

        elif ext == '.docx':
            import docx
            doc = docx.Document(tmp.name)
            text = '\n'.join(p.text for p in doc.paragraphs)

        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(tmp.name, read_only=True, data_only=True)
            try:
                ws = wb.active
                rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append('\t'.join(cells))
                text = '\n'.join(rows)
            finally:
                wb.close()

        else:  # .txt / .md
            text = content.decode('utf-8')

        return text
    finally:
        try:
            os.unlink(tmp.name)
        except PermissionError:
            pass  # Windows may still have locks; skip deletion


@app.post("/api/extract-file")
async def extract_file(file: UploadFile = File(...)):
    """Extract text from uploaded file (.pdf / .docx / .xlsx / .txt / .md)."""
    ext = os.path.splitext(file.filename)[1].lower()
    allowed = {'.pdf', '.docx', '.xlsx', '.xls', '.txt', '.md'}
    if ext not in allowed:
        return {"error": f"不支持的文件格式: {ext}，支持 PDF / Word / Excel / TXT / MD"}

    content = await file.read()

    try:
        text = _extract_text(content, ext)
    except Exception as e:
        return {"error": f"文件解析失败: {str(e)}"}

    if not text.strip():
        return {"error": "未能从文件中提取到文本内容"}

    return {"text": text}


@app.get("/api/health")
def health():
    return {"status": "ok"}
